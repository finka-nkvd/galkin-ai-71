from abc import ABC, abstractmethod
import json
import csv
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException



class DataProcessor(ABC):
    @abstractmethod
    def process_data(self, data):
        """Обработка данных"""
        pass


class CSVProcessor(DataProcessor):
    def __init__(self, text, file_name):
        self.text = text
        self.file_name = file_name

    def csv_input(self):
        file_exists = os.path.isfile(self.file_name)

        with open(self.file_name, 'a', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)

            if not file_exists:
                writer.writerow(['Text'])

            writer.writerow([self.text])

        print(f"Текст успешно добавлен в файл {self.file_name}")


class TXTProcessor(DataProcessor):
    def __init__(self, text, file_name):
        self.text = text
        self.file_name = file_name

    def txt_input(self):
        with open(self.file_name, 'a', encoding='utf-8') as txtfile:
            txtfile.write(self.text + '\n')

        print(f"Текст успешно добавлен в файл {self.file_name}")


class EXCELProcessor(DataProcessor):
    def __init__(self, text, file_name):
        self.text = text
        self.file_name = file_name

    def excel_input(self):
        sheet_name = "Sheet1"

        try:
            wb = load_workbook(self.file_name)
            sheet = wb[sheet_name]

            row_num = sheet.max_row + 1

            sheet.cell(row=row_num, column=1).value = self.text

            wb.save(self.file_name)

        except FileNotFoundError:
            df = pd.DataFrame([self.text], columns=['Text'])
            df.to_excel(self.file_name, sheet_name=sheet_name, index=False)

        except InvalidFileException:
            df = pd.DataFrame([self.text], columns=['Text'])
            df.to_excel(self.file_name, sheet_name=sheet_name, index=False)

        except KeyError:
            df = pd.DataFrame([self.text], columns=['Text'])
            df.to_excel(self.file_name, sheet_name=sheet_name, index=False)

        print(f"Текст успешно добавлен в файл {self.file_name} на лист {sheet_name}")


class JSONProcessor(DataProcessor):
    def __init__(self, key, text, file_name):
        self.text = text
        self.file_name = file_name
        self.key = key

    def json_input(self):
        new_data = {self.key: self.text}

        try:
            with open(self.file_name,  'r', encoding='utf-8') as jsonfile:
                data = json.load(jsonfile)
        except (FileNotFoundError, json.JSONDecodeError):
            data = {}

        data.update(new_data)

        with open(self.file_name, 'w', encoding='utf-8') as jsonfile:
            json.dump(data, jsonfile, indent=4, ensure_ascii=False)

        print(f"Данные успешно добавлены в файл {self.file_name}")


class Menu:
    def main(self):
        print('Введите тип данных для ввода:\n'
              '1. CSV\n'
              '2. TXT\n'
              '3. EXCEL\n'
              '4. JSON')
        file_type = input()

        if file_type == '1' or 'CSV' or 'csv':
            CSVProcessor(input('Введите текст для записи в файл: '), input("Введите название файла: "))
        elif file_type == '2' or 'TXT' or 'txt':
            TXTProcessor(input('Введите текст для записи в файл: '), input("Введите название файла: "))
        elif file_type == '3' or 'EXCEL' or 'excel':
            EXCELProcessor(input('Введите текст для записи в файл: '), input("Введите название файла: "))
        elif file_type == '4' or 'JSON' or 'json':
            JSONProcessor(input('Введите заголовок для записи в файл: '), input("Введите текст для записи в файл: "), input("Введите название файла: "))
        else:
            print('Неверный тип данных')

menu = Menu()
menu.main()
